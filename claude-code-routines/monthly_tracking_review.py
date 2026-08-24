"""Build the Monthly Tracking Review email from the Dialed tracking workbooks.

The tool is deliberately thin: it finds the month's workbooks, pulls the tracking-error
ratios, renders the ``CtP`` screenshots, and assembles an Outlook draft. The narrative
judgment is written by hand (or by the agent) and passed in with ``--narrative``.

Default behavior is a dry run. Use ``--draft`` to create an Outlook draft. ``--send``
exists but is deliberately separate; nothing is ever sent without it.

    python monthly_tracking_review.py readiness --month 2026-08
    python monthly_tracking_review.py facts     --month 2026-08
    python monthly_tracking_review.py draft     --month 2026-08 --narrative narrative.html

See ``MONTHLY_TRACKING_REVIEW.md`` for the methodology and the failure modes this guards
against -- in particular why a zero ``Proj`` must never be read as a slow model, and why the
``ChartObject``/``Chart.Export`` route for the screenshots is not used.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import shutil
import sys
import tempfile
import time
import uuid
from pathlib import Path

TRACKING_ROOT = Path(r"R:\QR\Resi_shared\tracking")
DIALED = TRACKING_ROOT / "Dialed"
UNDIALED = TRACKING_ROOT / "Undialed"

DEFAULT_TO = "LibreMax-Modeling"

# Cohort order is the order they appear in the email.
COHORTS = [
    ("STACR", re.compile(r"^tracking_STACR_.*_CRT_(\d{8})\.xlsx$", re.I)),
    ("CAS", re.compile(r"^tracking_CAS_.*_CRT_(\d{8})\.xlsx$", re.I)),
    ("NQM", re.compile(r"^tracking_(?!.*PSEUDO).*_NONQM_(\d{8})\.xlsx$", re.I)),
    ("JUMBO", re.compile(r"^tracking_.*_JUMBO_(\d{8})\.xlsx$", re.I)),
    ("HELOC", re.compile(r"^tracking_.*_HE_(\d{8})\.xlsx$", re.I)),
]

# Sheets whose ratios we report. CtP drives the narrative; CPR is quoted for HELOC/JUMBO.
METRIC_SHEETS = ("CtP", "CPR")


# --------------------------------------------------------------------------- discovery


def _scan(folder: Path) -> dict[str, list[tuple[dt.date, Path]]]:
    found: dict[str, list[tuple[dt.date, Path]]] = {name: [] for name, _ in COHORTS}
    if not folder.is_dir():
        return found
    for entry in folder.iterdir():
        if not entry.is_file():
            continue
        for name, pat in COHORTS:
            m = pat.match(entry.name)
            if m:
                try:
                    stamp = dt.datetime.strptime(m.group(1), "%Y%m%d").date()
                except ValueError:
                    break
                found[name].append((stamp, entry))
                break
    for name in found:
        found[name].sort(key=lambda t: (t[0], t[1].stat().st_mtime))
    return found


def resolve_month(month: str) -> tuple[int, int]:
    m = re.fullmatch(r"(\d{4})-(\d{2})", month)
    if not m:
        raise argparse.ArgumentTypeError(f"invalid --month {month!r}; expected YYYY-MM")
    return int(m.group(1)), int(m.group(2))


def pick_for_month(folder: Path, year: int, month: int) -> tuple[dict, dict]:
    """Latest file per cohort dated within the target month, plus what's missing."""
    scanned = _scan(folder)
    chosen, missing = {}, {}
    for name, _ in COHORTS:
        in_month = [(d, p) for d, p in scanned[name] if d.year == year and d.month == month]
        if in_month:
            chosen[name] = in_month[-1]
        else:
            newest = scanned[name][-1] if scanned[name] else None
            missing[name] = newest
    return chosen, missing


# ----------------------------------------------------------------------------- metrics


def _ratio_columns(ws) -> dict[str, int]:
    """window tag -> column index of its 'Ratio' sub-column"""
    out: dict[str, int] = {}
    for c in range(14, ws.max_column + 1):
        if ws.cell(row=3, column=c).value == "Ratio":
            for cc in range(c, 13, -1):
                tag = ws.cell(row=2, column=cc).value
                if isinstance(tag, str) and "Error" in tag:
                    out[tag.replace(" Error", "").strip()] = c
                    break
    return out


def _all_avg_rows(ws) -> list[int]:
    rows = []
    for r in range(4, ws.max_row + 1):
        if str(ws.cell(row=r, column=3).value or "").strip().upper() == "ALL AVG":
            rows.append(r)
    return rows


def _month_columns(ws) -> list[str]:
    out = []
    for c in range(14, ws.max_column + 1):
        v = ws.cell(row=2, column=c).value
        if hasattr(v, "year"):
            out.append(f"{v:%Y-%m}")
    return out


def _monthly_pairs(ws, row: int) -> list[tuple[str, object, object]]:
    """(month, Actual, Proj) for each month column at the given row."""
    pairs = []
    for c in range(14, ws.max_column + 1):
        hdr = ws.cell(row=2, column=c).value
        if hasattr(hdr, "year"):
            pairs.append((
                f"{hdr:%Y-%m}",
                ws.cell(row=row, column=c).value,
                ws.cell(row=row, column=c + 1).value,
            ))
    return pairs


def _proj_gaps(pairs) -> list[str]:
    """Months whose projection is missing or zero.

    A zero Proj is not a slow model -- it is an absent projection, and because the
    window ratios are sum(Proj)/sum(Actual) it silently drags every window that
    contains it. HELOC 2026-06 did exactly this and understated the 3M ratio by
    ~0.32. Always surface these before quoting a ratio.
    """
    return [m for m, _a, p in pairs if not isinstance(p, (int, float)) or p == 0]


def extract_metrics(path: Path) -> dict:
    """Ratios at the WAC 'ALL AVG' row -- the row that reproduces the numbers quoted
    in prior reviews. Also keeps the AGE row for cross-checking."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    out: dict = {"file": str(path), "sheets": {}}
    for sheet in METRIC_SHEETS:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        cols = _ratio_columns(ws)
        rows = _all_avg_rows(ws)
        if not cols or not rows:
            continue
        entry = {"months": _month_columns(ws), "windows": sorted(cols)}
        for label, idx in (("age", 0), ("wac", 1)):
            if idx < len(rows):
                r = rows[idx]
                entry[label] = {
                    w: (round(v, 3) if isinstance(v, (int, float)) else None)
                    for w, c in cols.items()
                    for v in (ws.cell(row=r, column=c).value,)
                }
                entry[f"{label}_row"] = r

        # gap detection on the row we actually quote
        if len(rows) > 1:
            pairs = _monthly_pairs(ws, rows[1])
            gaps = _proj_gaps(pairs)
            entry["monthly"] = [
                {"month": m,
                 "actual": round(a, 4) if isinstance(a, (int, float)) else None,
                 "proj": round(p, 4) if isinstance(p, (int, float)) else None}
                for m, a, p in pairs
            ]
            entry["proj_gap_months"] = gaps
            if gaps:
                clean = [(m, a, p) for m, a, p in pairs[:3]
                         if isinstance(p, (int, float)) and p != 0]
                aa = [a for _, a, _ in clean if isinstance(a, (int, float))]
                pp = [p for _, _, p in clean]
                entry["ratio_3m_excl_gaps"] = (
                    round(sum(pp) / sum(aa), 3) if aa and sum(aa) else None
                )
                entry["ratio_3m_excl_gaps_n"] = len(clean)
        out["sheets"][sheet] = entry
    wb.close()
    return out


def gap_warnings(report: dict) -> list[str]:
    """Human-readable warnings for every cohort/sheet with a missing projection."""
    warnings = []
    for coh, data in report.get("cohorts", {}).items():
        for sheet, entry in data.get("sheets", {}).items():
            gaps = entry.get("proj_gap_months") or []
            if not gaps:
                continue
            adj = entry.get("ratio_3m_excl_gaps")
            n = entry.get("ratio_3m_excl_gaps_n")
            msg = (f"{coh} {sheet}: projection missing/zero for {', '.join(gaps)} -- "
                   f"every window ratio containing it is understated")
            if adj is not None:
                msg += f"; 3M ratio excluding it = {adj} (over {n} of 3 months)"
            warnings.append(msg)
    return warnings


def dial_status(name: str, stamp: dt.date, dialed: Path) -> str:
    """'no dial' if the Undialed twin is numerically identical, else 'dial applied'.
    Returns 'unverified' when no Undialed file exists for that date."""
    twin = UNDIALED / dialed.name
    if not twin.is_file():
        return "unverified (no Undialed file for this date)"
    try:
        import openpyxl

        out = []
        for p in (dialed, twin):
            wb = openpyxl.load_workbook(p, data_only=True)
            ws = wb["CtP"]
            rows = _all_avg_rows(ws)
            cols = _ratio_columns(ws)
            r = rows[1] if len(rows) > 1 else rows[0]
            out.append(tuple(ws.cell(row=r, column=c).value for c in sorted(cols.values())))
            wb.close()
        return "no dial applied" if out[0] == out[1] else "dial applied"
    except Exception as exc:  # pragma: no cover
        return f"unverified ({exc})"


# ------------------------------------------------------------------------ chart render


def _looks_blank(path: Path) -> bool:
    from PIL import Image

    try:
        im = Image.open(path).convert("L")
    except Exception:
        return True
    colors = im.getcolors(maxcolors=1 << 24) or []
    if not colors:
        return True
    total = sum(c for c, _ in colors)
    top = max(c for c, _ in colors)
    return (top / total) > 0.995


def render_ctp_images(chosen: dict, workdir: Path) -> dict[str, Path]:
    """Screenshot each cohort's CtP sheet, rows 1..(2nd 'ALL AVG'), columns A:AE --
    the range used in prior reviews (AGE + WAC blocks)."""
    import pythoncom
    import win32com.client
    from PIL import Image, ImageGrab

    workdir.mkdir(parents=True, exist_ok=True)
    images: dict[str, Path] = {}

    xl = win32com.client.Dispatch("Excel.Application")
    xl.Visible = True  # CopyPicture is unreliable without a real window
    xl.DisplayAlerts = False
    try:
        for name, (stamp, src) in chosen.items():
            # never open the shared original for anything but a read-only copy
            local = workdir / src.name
            if not local.exists():
                shutil.copy2(src, local)

            wb = xl.Workbooks.Open(str(local), UpdateLinks=0)
            try:
                ws = wb.Worksheets("CtP")
                ws.Activate()
                hits = []
                for r in range(1, 400):
                    v = ws.Cells(r, 3).Value
                    if v and str(v).strip().upper() == "ALL AVG":
                        hits.append(r)
                    if len(hits) >= 2:
                        break
                if not hits:
                    raise RuntimeError(f"{name}: no 'ALL AVG' row found in CtP")
                last = hits[1] if len(hits) >= 2 else hits[0]
                rng = ws.Range(f"A1:AE{last}")

                grabbed = None
                for _ in range(5):
                    rng.CopyPicture(Appearance=1, Format=2)  # xlScreen, xlBitmap
                    for _ in range(20):
                        pythoncom.PumpWaitingMessages()
                        time.sleep(0.25)
                        img = ImageGrab.grabclipboard()
                        if isinstance(img, Image.Image):
                            grabbed = img
                            break
                    if grabbed is not None:
                        break
                if grabbed is None:
                    raise RuntimeError(f"{name}: clipboard never yielded a bitmap")

                dest = workdir / f"ctp_{name}.png"
                grabbed.save(dest, "PNG")
                if _looks_blank(dest):
                    raise RuntimeError(f"{name}: rendered image is blank ({dest})")
                images[name] = dest
                print(f"  [OK] {name}: A1:AE{last} -> {dest.name} ({dest.stat().st_size:,} b)")
            finally:
                wb.Close(SaveChanges=False)
    finally:
        xl.Quit()
    return images


# -------------------------------------------------------------------------- assembly


def build_html(narrative_html: str, images: dict[str, Path], cids: dict[str, str]) -> str:
    blocks = [narrative_html]
    for name in [n for n, _ in COHORTS if n in images]:
        blocks.append(
            f'<p style="margin:14pt 0 4pt 0"><b>{name} CtoP</b></p>'
            f'<p><img src="cid:{cids[name]}" style="max-width:100%"></p>'
        )
    blocks.append(
        '<p style="margin-top:16pt">Best,<br>Howard Zeng<br>QR</p>'
    )
    return (
        '<html><body style="font-family:Calibri,sans-serif;font-size:11pt;color:#000">'
        + "".join(blocks)
        + "</body></html>"
    )


def create_draft(subject: str, to: str, html: str, images: dict[str, Path],
                 cids: dict[str, str], attachments: list[Path], send: bool) -> None:
    import win32com.client

    outlook = win32com.client.Dispatch("Outlook.Application")
    mail = outlook.CreateItem(0)
    mail.To = to
    mail.Subject = subject
    mail.HTMLBody = html

    for path in attachments:
        mail.Attachments.Add(str(path))

    for name, path in images.items():
        att = mail.Attachments.Add(str(path))
        try:
            att.PropertyAccessor.SetProperty(
                "http://schemas.microsoft.com/mapi/proptag/0x3712001F", cids[name]
            )
        except Exception as exc:
            print(f"  [WARN] could not set content-id for {name}: {exc}")

    if send:
        mail.Send()
        print(f"[OK] sent to {to}")
    else:
        mail.Save()
        print(f"[OK] draft created for {to} -- review in Outlook Drafts before sending")


# ------------------------------------------------------------------------------- cli


def cmd_facts(args) -> int:
    year, month = resolve_month(args.month)
    chosen, missing = pick_for_month(DIALED, year, month)

    report = {"month": args.month, "source": str(DIALED), "cohorts": {}, "missing": {}}
    for name, (stamp, path) in chosen.items():
        report["cohorts"][name] = {
            "as_of": stamp.isoformat(),
            "dial": dial_status(name, stamp, path),
            **extract_metrics(path),
        }
    for name, newest in missing.items():
        report["missing"][name] = (
            {"newest_available": newest[0].isoformat(), "file": newest[1].name}
            if newest else None
        )

    print(json.dumps(report, indent=2))
    if missing:
        print(f"\n[NOTE] no {args.month} file for: {', '.join(sorted(missing))}",
              file=sys.stderr)
    for w in gap_warnings(report):
        print(f"[DATA GAP] {w}", file=sys.stderr)
    return 0


def cmd_readiness(args) -> int:
    """Pre-flight: is the month's data actually reviewable yet?

    Answers the question that is easy to miss until you sit down to write the review --
    which cohorts have landed, and are any of the ones that landed unusable. Exits 4 when
    something is missing or gapped so it can gate a later step.
    """
    year, month = resolve_month(args.month)
    chosen, missing = pick_for_month(DIALED, year, month)

    gaps: dict[str, list[str]] = {}
    for name, (_stamp, path) in chosen.items():
        metrics = extract_metrics(path)
        found = []
        for sheet, entry in metrics.get("sheets", {}).items():
            for m in entry.get("proj_gap_months") or []:
                found.append(f"{sheet} {m}")
        if found:
            gaps[name] = found

    print(f"Monthly tracking readiness -- {args.month}")
    print(f"source: {DIALED}\n")
    print(f"  {'cohort':<8} {'as-of':<12} status")
    print(f"  {'-' * 8} {'-' * 12} {'-' * 52}")
    for name, _pat in COHORTS:
        if name in chosen:
            stamp, _path = chosen[name]
            if name in gaps:
                status = f"DATA GAP: projection zero for {', '.join(gaps[name])}"
            else:
                status = "ready"
            print(f"  {name:<8} {stamp.isoformat():<12} {status}")
        else:
            newest = missing.get(name)
            detail = (f"MISSING (newest on file {newest[0].isoformat()})"
                      if newest else "MISSING (no file ever found)")
            print(f"  {name:<8} {'--':<12} {detail}")

    total = len(COHORTS)
    print(f"\n  {len(chosen)} of {total} cohorts present"
          + (f"; {len(missing)} missing" if missing else "")
          + (f"; {len(gaps)} with a data gap" if gaps else ""))

    if not missing and not gaps:
        print("\n  READY -- the review can be written from this month's data.")
        return 0

    print("\n  NOT READY")
    if missing:
        print(f"    - regenerate tracking for: {', '.join(sorted(missing))}")
        print("    - a missing cohort usually means its tracking pipeline failed; check the")
        print("      Jenkins monitor for the quant-Monthly-ResiTracking jobs")
    if gaps:
        print(f"    - repopulate the missing projections for: {', '.join(sorted(gaps))}")
        print("    - until then, every tracking-error window containing that month is")
        print("      understated and must not be quoted at face value")
    return 4


def cmd_draft(args) -> int:
    year, month = resolve_month(args.month)
    chosen, missing = pick_for_month(DIALED, year, month)
    if not chosen:
        print(f"[FAIL] no tracking files at all for {args.month} in {DIALED}")
        return 2

    print(f"cohorts for {args.month}: {', '.join(sorted(chosen))}")
    if missing:
        print(f"missing: {', '.join(sorted(missing))}")

    # Surface projection gaps before anything is written into an email. A zero Proj
    # understates every window ratio that contains it, so the narrative must account
    # for it rather than read the ratio at face value.
    probe = {"cohorts": {n: extract_metrics(p) for n, (_d, p) in chosen.items()}}
    warnings = gap_warnings(probe)
    for w in warnings:
        print(f"[DATA GAP] {w}")
    if warnings and not args.ack_gaps:
        print("\n[FAIL] data gaps present. Confirm the narrative accounts for them, "
              "then re-run with --ack-gaps.")
        return 3

    narrative = Path(args.narrative).read_text(encoding="utf-8").strip()
    if not narrative:
        print("[FAIL] --narrative file is empty")
        return 2

    workdir = Path(args.workdir) if args.workdir else Path(tempfile.mkdtemp(prefix="mtr_"))
    print(f"workdir: {workdir}")

    print("rendering CtP images...")
    images = render_ctp_images(chosen, workdir)

    cids = {name: f"ctp-{name.lower()}-{uuid.uuid4()}" for name in images}
    html = build_html(narrative, images, cids)

    label = dt.date(year, month, 1).strftime("%B")
    subject = args.subject or f"Monthly Tracking Review - {label}"
    attachments = [p for _, p in sorted(chosen.values(), key=lambda t: t[0])]

    if args.dry_run or not (args.draft or args.send):
        out = workdir / "preview.html"
        out.write_text(html, encoding="utf-8")
        print(f"[DRY RUN] subject : {subject}")
        print(f"[DRY RUN] to      : {args.to}")
        print(f"[DRY RUN] attach  : {[p.name for p in attachments]}")
        print(f"[DRY RUN] images  : {[p.name for p in images.values()]}")
        print(f"[DRY RUN] preview : {out}")
        return 0

    create_draft(subject, args.to, html, images, cids, attachments, send=args.send)
    return 0


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)

    default_month = dt.date.today().strftime("%Y-%m")

    r = sub.add_parser("readiness",
                       help="which cohorts have landed and are usable (exit 4 if not)")
    r.add_argument("--month", default=default_month, help="YYYY-MM (default: this month)")
    r.set_defaults(func=cmd_readiness)

    f = sub.add_parser("facts", help="print the month's tracking ratios as JSON")
    f.add_argument("--month", default=default_month, help="YYYY-MM (default: this month)")
    f.set_defaults(func=cmd_facts)

    d = sub.add_parser("draft", help="render charts and assemble the email")
    d.add_argument("--month", default=default_month, help="YYYY-MM (default: this month)")
    d.add_argument("--narrative", required=True,
                   help="HTML fragment holding Summary / Next Steps prose")
    d.add_argument("--to", default=DEFAULT_TO)
    d.add_argument("--subject", default=None)
    d.add_argument("--workdir", default=None, help="where copies and PNGs are written")
    d.add_argument("--ack-gaps", action="store_true",
                   help="proceed despite missing projections (narrative must address them)")
    mode = d.add_mutually_exclusive_group()
    mode.add_argument("--dry-run", action="store_true", help="write preview.html only")
    mode.add_argument("--draft", action="store_true", help="create an Outlook draft")
    mode.add_argument("--send", action="store_true", help="send immediately")
    d.set_defaults(func=cmd_draft)

    return p


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
