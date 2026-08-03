from __future__ import annotations

import importlib.util
import sys
from pathlib import Path
from types import ModuleType, SimpleNamespace

import pytest
from openpyxl import load_workbook

_REPO_ROOT = Path(__file__).resolve().parents[1]
_ROLL_RATE_ROOT = _REPO_ROOT / "roll-rate-model"
_ROLL_RATE_PYTHON = (_ROLL_RATE_ROOT / "python").resolve()


def _import_vendored_simengine() -> tuple[ModuleType, ModuleType, ModuleType]:
    original_sys_path = sys.path.copy()
    try:
        sys.path.insert(0, str(_ROLL_RATE_PYTHON))
        import simengine.data_prep as data_prep
        import simengine.register_vars as register_vars
        import simengine.runner as runner

        for module in (data_prep, register_vars, runner):
            expected_path = _ROLL_RATE_PYTHON.joinpath(
                *module.__name__.split(".")
            ).with_suffix(".py")
            assert Path(module.__file__).samefile(expected_path)
        return data_prep, register_vars, runner
    finally:
        sys.path[:] = original_sys_path


data_prep, register_vars, runner = _import_vendored_simengine()

_UPDATE_MACRO_PATH = _ROLL_RATE_ROOT / "input" / "macro" / "update_macro.py"
_UPDATE_MACRO_SPEC = importlib.util.spec_from_file_location(
    "roll_rate_update_macro",
    _UPDATE_MACRO_PATH,
)
update_macro = importlib.util.module_from_spec(_UPDATE_MACRO_SPEC)
_UPDATE_MACRO_SPEC.loader.exec_module(update_macro)


def test_run_simulation_explicit_n_per_overrides_config(
    tmp_path: Path,
) -> None:
    config = {
        "n_per": 4,
        "status_to_roll": {"C": ["C"]},
    }
    loan = {
        "loan_id": "smoke",
        "end_bal": 100.0,
        "int_rate": 0.12,
        "term": 12,
        "loan_age": 0,
        "status": "C",
        "r_dt": "2026-01-31",
    }

    result = runner.run_simulation(
        [loan],
        str(tmp_path),
        n_per=3,
        config=config,
        mode="sequential",
    )

    assert result["errors"] == []
    assert result["dm"].n_per == 3
    assert result["dm"].status_to_roll == {"C": ["C"]}
    assert len(result["loan_results"]) == 1
    assert len(result["loan_results"][0]["cf"]) == 3
    assert config["n_per"] == 4


def test_run_simulation_omitted_n_per_uses_config_horizon(
    tmp_path: Path,
) -> None:
    config = {
        "n_per": 2,
        "status_to_roll": {"C": ["C"]},
    }
    loan = {
        "loan_id": "config-horizon",
        "end_bal": 100.0,
        "int_rate": 0.12,
        "term": 12,
        "loan_age": 0,
        "status": "C",
        "r_dt": "2026-01-31",
    }

    result = runner.run_simulation(
        [loan],
        str(tmp_path),
        config=config,
        mode="sequential",
    )

    assert result["errors"] == []
    assert result["dm"].n_per == 2
    assert len(result["cf_sum"]) == 2
    assert len(result["loan_results"]) == 1
    assert len(result["loan_results"][0]["cf"]) == 2


def test_write_results_xlsx_includes_period_indexed_grouped_metrics(
    tmp_path: Path,
) -> None:
    loan = {
        "loan_id": "seasoned-stay",
        "end_bal": 100.0,
        "int_rate": 0.12,
        "term": 12,
        "grade": "A",
        "loan_age": 7,
        "status": "C",
        "r_dt": "2026-01-31",
    }
    result = runner.run_simulation(
        [loan],
        str(tmp_path),
        n_per=2,
        status_to_roll={"C": ["C"]},
        mode="sequential",
    )
    output_path = tmp_path / "sim_results.xlsx"

    runner.write_results_xlsx(
        result,
        str(output_path),
        group_by=["term", "grade"],
    )

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    assert {
        "Portfolio",
        "Metrics_Portfolio",
        "Metrics_Grouped",
        "Metrics_Grouped_Period",
    } <= set(workbook.sheetnames)

    worksheet = workbook["Metrics_Grouped_Period"]
    rows = list(worksheet.iter_rows(values_only=True))
    header = list(rows[0])
    assert header[:3] == ["term", "grade", "period"]
    assert "fromC_stay" in header

    data_rows = rows[1:]
    assert [row[2] for row in data_rows] == [1, 2]
    assert {(str(row[0]), row[1]) for row in data_rows} == {("12", "A")}
    stay_index = header.index("fromC_stay")
    assert [row[stay_index] for row in data_rows] == [1.0, 1.0]
    workbook.close()


def test_run_simulation_empty_config_preserves_explicit_status_to_roll(
    tmp_path: Path,
) -> None:
    config = {}
    loan = {
        "loan_id": "explicit-status-map",
        "end_bal": 100.0,
        "int_rate": 0.12,
        "term": 12,
        "loan_age": 0,
        "status": "C",
        "r_dt": "2026-01-31",
    }

    result = runner.run_simulation(
        [loan],
        str(tmp_path),
        n_per=2,
        status_to_roll={"C": ["C"]},
        config=config,
        mode="sequential",
    )

    assert result["errors"] == []
    assert result["dm"].n_per == 2
    assert result["dm"].status_to_roll == {"C": ["C"]}
    assert len(result["loan_results"]) == 1
    assert len(result["loan_results"][0]["cf"]) == 2
    assert config == {}


def test_gam_term_stack_maps_to_supported_terms() -> None:
    terms = [24, 36, 48, 60, 84, None]

    assert [register_vars._gam_term_stack(term) for term in terms] == [
        36,
        36,
        60,
        60,
        60,
        None,
    ]


def test_load_cpi_lookup_allows_zero_extension(tmp_path: Path) -> None:
    cpi_path = tmp_path / "CPIAUCNS.csv"
    cpi_path.write_text(
        "DATE,CPIAUCNS\n"
        "1/1/2024,308.417\n"
        "2024-02-01,310.326\n",
        encoding="utf-8",
    )

    assert data_prep.load_cpi_lookup(str(cpi_path), extend_months=0) == {
        "2024-01": 308.417,
        "2024-02": 310.326,
    }


def test_update_cpi_macro_fields_recomputes_active_inflators() -> None:
    loan = {"cpi_inflator_12": -1.0, "cpi_inflator_36": -1.0}
    cpi_lookup = {"2024-01": 310.368, "2023-01": 300.0, "2021-01": 250.0}

    runner._update_cpi_macro_fields(
        loan,
        cpi_lookup,
        {"cpi_inflator_12", "cpi_inflator_36"},
        "2024-01",
    )

    assert loan["cpi_inflator_12"] == 0.0346
    assert loan["cpi_inflator_36"] == 0.2415


def test_update_cpi_macro_fields_leaves_inactive_inflator_unchanged() -> None:
    loan = {"cpi_inflator_12": 0.12, "cpi_inflator_36": 0.36}
    cpi_lookup = {"2024-01": 120.0, "2023-01": 100.0, "2021-01": 80.0}

    runner._update_cpi_macro_fields(
        loan,
        cpi_lookup,
        {"cpi_inflator_12"},
        "2024-01",
    )

    assert loan == {"cpi_inflator_12": 0.2, "cpi_inflator_36": 0.36}


@pytest.mark.parametrize(
    ("cpi_lookup", "expected"),
    [
        pytest.param(
            {"2023-01": 100.0, "2021-01": 80.0},
            {"cpi_inflator_12": 0.12, "cpi_inflator_36": 0.36},
            id="current-cpi-missing",
        ),
        pytest.param(
            {"2024-01": 120.0, "2023-01": 100.0},
            {"cpi_inflator_12": 0.2, "cpi_inflator_36": 0.36},
            id="36-month-lag-missing",
        ),
        pytest.param(
            {"2024-01": 120.0, "2021-01": 80.0},
            {"cpi_inflator_12": 0.12, "cpi_inflator_36": 0.5},
            id="12-month-lag-missing",
        ),
    ],
)
def test_update_cpi_macro_fields_freezes_missing_values(
    cpi_lookup: dict[str, float],
    expected: dict[str, float],
) -> None:
    loan = {"cpi_inflator_12": 0.12, "cpi_inflator_36": 0.36}

    runner._update_cpi_macro_fields(
        loan,
        cpi_lookup,
        {"cpi_inflator_12", "cpi_inflator_36"},
        "2024-01",
    )

    assert loan == expected


def test_fetch_fred_csv_normalizes_dates_to_year_month(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class FakeResponse:
        text = (
            "observation_date,CPIAUCNS\n"
            "2024-01-01,308.417\n"
            "2024-02-01,.\n"
            "2024-11-01,315.664\n"
        )

        def raise_for_status(self) -> None:
            pass

    fake_requests = SimpleNamespace(
        get=lambda url, timeout: FakeResponse(),
    )
    monkeypatch.setitem(sys.modules, "requests", fake_requests)

    assert update_macro.fetch_fred_csv("CPIAUCNS") == [
        ("2024-01", 308.417),
        ("2024-11", 315.664),
    ]
