"""Convert a JPM BondStudio cashflow export into the long format plot_cdr_compare reads.

The export is a wide sheet: repeating 7-column blocks (Date, Balance, Principal,
Interest, CPR, CDR, Severity), one block per bond x scenario. Row 5 carries the
block label, row 6 the column names, row 7 onward the data -- and row 7 is the
as-of row with zero flows, so the first projection month is row 8.

Deal and scenario names are translated to ours so the JPM series lands in the same
panel as prod and dialed. JPM's Base Case is not our Base (ours is the GS MSA
forecast), so the base lines will not tie out; the comparison is shape, not level.

Usage::

    python emailer/jpm_cashflow_to_long.py "C:\\Users\\hzeng\\Downloads\\Cashflows-...xlsx"
    python emailer/plot_cdr_compare.py --csv <merged csv> --deals "<the six>"
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook

EMAILER_DIR = Path(__file__).resolve().parent
DEFAULT_BASE_CSV = Path(r"S:\QR\hzeng\heloc_jumbo_cdr_compare.csv")
DEFAULT_OUT_CSV = Path(r"S:\QR\hzeng\heloc_jumbo_cdr_compare_with_jpm.csv")

BLOCK_WIDTH = 7
FIRST_DATA_ROW = 8  # row 7 is the as-of row (zero flows); month 1 is row 8

# JPM shorthand -> our bbg_deal_name
DEAL_MAP = {
    "ACHM-24HE1": "ACHM 2024-HE1",
    "VSTA-4CES1": "VSTA 2024-CES1",
    "JPMMT-26HE1": "JPMMT 2026-HE1",
    "JPMMT-6LTV1": "JPMMT 2026-LTV1",
    "MSRM-243": "MSRM 2024-3",
    "JPMMT-195": "JPMMT 2019-5",
}
# JPM scenario -> ours. SuperBull/SuperBear are HPI-only on our side (2y HPA base
# +12% / -35%), so the HPI ladder is the right analogue, not the rate scenarios.
SCEN_MAP = {
    "Base Case": "Base",
    "HPI Bullish Recovery": "SuperBull",
    "HPI Severely Negative": "SuperBear",
}
PURPOSE_NAME = "JPM"
COLUMNS = ["purpose", "purpose_name", "bbg_name", "poolid", "scen", "month",
           "sbal", "cpr", "cdr", "sev", "d60", "wac"]


def _build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="JPM cashflow export -> long-format CSV.")
    p.add_argument("xlsx", type=Path, help="the BondStudio Cashflows-*.xlsx")
    p.add_argument("--base-csv", type=Path, default=DEFAULT_BASE_CSV,
                   help="existing prod/dialed CSV to merge with (omit with --jpm-only)")
    p.add_argument("--out", type=Path, default=DEFAULT_OUT_CSV)
    p.add_argument("--jpm-only", action="store_true", help="write only the JPM rows")
    p.add_argument("--max-month", type=int, default=120)
    return p


def parse_label(label: str) -> tuple[str, str] | None:
    """'ACHM-24HE1 B (B) - Base Case (Forward)' -> ('ACHM 2024-HE1', 'Base')."""
    if not label or " - " not in label:
        return None
    left, right = label.split(" - ", 1)
    shorthand = left.split()[0]
    scen = right.replace("(Forward)", "").strip()
    deal, our_scen = DEAL_MAP.get(shorthand), SCEN_MAP.get(scen)
    return (deal, our_scen) if deal and our_scen else None


def extract(xlsx: Path, max_month: int) -> list[dict[str, object]]:
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb.active
    grid = list(ws.iter_rows(values_only=True))
    labels = grid[4]
    out: list[dict[str, object]] = []
    skipped: list[str] = []

    for block in range(ws.max_column // BLOCK_WIDTH):
        c0 = block * BLOCK_WIDTH
        parsed = parse_label(str(labels[c0]) if c0 < len(labels) and labels[c0] else "")
        if parsed is None:
            if c0 < len(labels) and labels[c0]:
                skipped.append(str(labels[c0]))
            continue
        deal, scen = parsed
        month = 0
        for row in grid[FIRST_DATA_ROW - 1:]:
            if c0 >= len(row) or row[c0] is None:
                break
            month += 1
            if month > max_month:
                break
            get = lambda off: row[c0 + off] if c0 + off < len(row) else None
            out.append({
                "purpose": PURPOSE_NAME,
                "purpose_name": PURPOSE_NAME,
                "bbg_name": deal,
                "poolid": f"JPM_{deal.replace(' ', '_')}",
                "scen": scen,
                "month": month,
                "sbal": get(1),   # Balance
                "cpr": get(4),
                "cdr": get(5),
                "sev": get(6),
                "d60": "",
                "wac": "",
            })
    if skipped:
        print("  [warn] no mapping, skipped: %s" % "; ".join(sorted(set(skipped))))
    return out


def main(argv: list[str] | None = None) -> int:
    args = _build_parser().parse_args(argv)
    rows = extract(args.xlsx, args.max_month)
    if not rows:
        raise SystemExit("  no JPM rows parsed -- check DEAL_MAP / SCEN_MAP against row 5")

    deals = sorted({str(r["bbg_name"]) for r in rows})
    scens = sorted({str(r["scen"]) for r in rows})
    print(f"  parsed {len(rows)} JPM rows: {len(deals)} deals x {len(scens)} scenarios")
    print(f"    deals     {', '.join(deals)}")
    print(f"    scenarios {', '.join(scens)}")

    args.out.parent.mkdir(parents=True, exist_ok=True)
    with args.out.open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNS)
        w.writeheader()
        if not args.jpm_only:
            with args.base_csv.open(encoding="utf-8", newline="") as src:
                kept = 0
                for r in csv.DictReader(src):
                    w.writerow({c: r.get(c, "") for c in COLUMNS})
                    kept += 1
            print(f"  carried {kept} prod/dialed rows from {args.base_csv.name}")
        w.writerows(rows)
    print(f"  wrote {args.out}")
    print("  now:  python emailer/plot_cdr_compare.py --csv %s --deals \"%s\""
          % (args.out, ",".join(deals)))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
