"""Generate NQM vector comparison plots for email updates."""

from __future__ import annotations

import argparse
import re
from datetime import datetime
from pathlib import Path

from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure
from openpyxl import load_workbook

EMAILER_DIR = Path(__file__).resolve().parent
DEFAULT_WORKBOOK = Path(r"S:\QR\hzeng\vector_compare_nqm_v7.xlsx")
DEFAULT_SHEET = "vector_compare_0925"
DEFAULT_OUTPUT_DIR = EMAILER_DIR / "assets" / "nqm_vector_compare"
DEFAULT_DEALS = [
    "CROSS 2025-H7",
    "NYMT 2026-INV1",
    "BARC 2026-NQM1",
    "ADMT 2025-NQM1",
    "GCAT 2025-NQM3",
    "CHNGE 2023-4",
    "OBX 2026-NQM1",
]
DEFAULT_SCENARIOS = ["Base", "ParallelDn200", "ParallelDn300", "ParallelUp200"]
MODEL_STYLES = {
    "Prod": {"label": "Prod", "color": "#1f77b4", "linewidth": 2.2},
    "New NonQM": {"label": "New NQM", "color": "#f28c28", "linewidth": 2.4},
    "JPM": {"label": "JPM", "color": "#808080", "linewidth": 2.0},
    "New NonQM V9 No Decay Dialed": {"label": "Dialed NQM", "color": "#d4a017", "linewidth": 2.4},
}
SCENARIO_LABELS = {
    "Base": "Base",
    "ParallelDn200": "RateDown 200",
    "ParallelDn300": "RateDown 300",
    "ParallelUp200": "RateUp 200",
    "SuperBear": "SuperBear",
    "SuperBull": "SuperBull",
}


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Plot NQM vector comparisons by deal and scenario.")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK, help=f"Input workbook. Default: {DEFAULT_WORKBOOK}")
    parser.add_argument("--sheet", default=DEFAULT_SHEET, help=f"Input sheet. Default: {DEFAULT_SHEET}")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR, help=f"Output image folder. Default: {DEFAULT_OUTPUT_DIR}")
    parser.add_argument("--deals", nargs="*", default=DEFAULT_DEALS, help="BBG Deal names to plot.")
    parser.add_argument("--scenarios", nargs="*", default=DEFAULT_SCENARIOS, help="Scenarios to plot.")
    parser.add_argument("--metric", choices=["CPR", "CDR"], default="CPR", help="Vector metric to plot. Default: CPR.")
    parser.add_argument("--max-month", type=int, default=120, help="Maximum projection month to show. Default: 120.")
    parser.add_argument("--exclude-no-decay", action="store_true", help="Do not plot New NonQM No Decay.")
    return parser


def _slug(value: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "_", value.strip()).strip("_").lower()
    return slug or "plot"


def _as_float(value: object) -> float | None:
    try:
        if value is None or value == "":
            return None
        return float(value)
    except (TypeError, ValueError):
        return None


def _month_diff(start: datetime, current: datetime) -> int:
    return (current.year - start.year) * 12 + current.month - start.month + 1


def _load_data(workbook: Path, sheet: str, metric: str) -> list[dict[str, object]]:
    wb = load_workbook(workbook, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        required = ["BBG Deal", "Model", "Scenario", "Date", "Month", metric]
        missing = [column for column in required if column not in headers]
        if missing:
            raise ValueError(f"Sheet {sheet!r} missing required columns: {missing}")
        idx = {column: headers.index(column) for column in required}

        records: list[dict[str, object]] = []
        for row in rows:
            deal = row[idx["BBG Deal"]]
            model = row[idx["Model"]]
            scenario = row[idx["Scenario"]]
            date = row[idx["Date"]]
            month = _as_float(row[idx["Month"]])
            value = _as_float(row[idx[metric]])
            if not deal or not model or not scenario or value is None:
                continue
            records.append(
                {
                    "deal": str(deal).strip(),
                    "model": str(model).strip(),
                    "scenario": str(scenario).strip(),
                    "month": month,
                    "date": date,
                    "value": value,
                }
            )
        return _fill_missing_months(records)
    finally:
        wb.close()


def _fill_missing_months(records: list[dict[str, object]]) -> list[dict[str, object]]:
    group_start_dates: dict[tuple[str, str, str], datetime] = {}
    for record in records:
        if record["month"] is not None:
            continue
        date = record["date"]
        if not isinstance(date, datetime):
            continue
        key = (str(record["deal"]), str(record["model"]), str(record["scenario"]))
        current_start = group_start_dates.get(key)
        if current_start is None or date < current_start:
            group_start_dates[key] = date

    for record in records:
        if record["month"] is not None:
            continue
        date = record["date"]
        if not isinstance(date, datetime):
            continue
        key = (str(record["deal"]), str(record["model"]), str(record["scenario"]))
        start_date = group_start_dates.get(key)
        if start_date is not None:
            record["month"] = float(_month_diff(start_date, date))

    return [record for record in records if record["month"] is not None]


def _models_to_plot(exclude_no_decay: bool) -> dict[str, dict[str, object]]:
    styles = dict(MODEL_STYLES)
    if exclude_no_decay:
        styles.pop("New NonQM No Decay", None)
        styles.pop("New NonQM No Decay Dial", None)
        styles.pop("New NonQM V9 No Decay Dialed", None)
    return styles


def plot_deal_scenario(
    records: list[dict[str, object]],
    deal: str,
    scenario: str,
    output_dir: Path,
    max_month: int,
    model_styles: dict[str, dict[str, object]],
    metric: str,
) -> Path | None:
    subset = [
        record
        for record in records
        if str(record["deal"]).upper() == deal.upper()
        and record["scenario"] == scenario
        and float(record["month"]) <= max_month
        and str(record["model"]) in model_styles
    ]
    if not subset:
        return None

    fig = Figure(figsize=(8.6, 4.8), dpi=150)
    FigureCanvasAgg(fig)
    ax = fig.subplots()
    for model, style in model_styles.items():
        model_rows = sorted(
            [record for record in subset if record["model"] == model],
            key=lambda record: float(record["month"]),
        )
        if not model_rows:
            continue
        ax.plot(
            [float(record["month"]) for record in model_rows],
            [float(record["value"]) for record in model_rows],
            label=str(style["label"]),
            color=str(style["color"]),
            linewidth=float(style["linewidth"]),
        )

    scenario_label = SCENARIO_LABELS.get(scenario, scenario)
    ax.set_title(f"{deal} - {scenario_label}", fontsize=13, fontweight="bold")
    ax.set_xlabel("Projection Month")
    ax.set_ylabel(f"{metric} (%)")
    ax.grid(True, alpha=0.25)
    ax.set_xlim(left=1, right=max_month)
    values = sorted(float(record["value"]) for record in subset)
    if values:
        q99 = values[min(int(len(values) * 0.99), len(values) - 1)]
        ax.set_ylim(bottom=0, top=max(q99 * 1.15, max(values) * 1.03))
    ax.legend(loc="upper right", frameon=True)
    fig.tight_layout()

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{_slug(deal)}_{_slug(metric)}_{_slug(scenario)}.png"
    fig.savefig(output_path, bbox_inches="tight")
    return output_path


def main(argv: list[str] | None = None) -> int:
    args = _build_parser().parse_args(argv)
    records = _load_data(args.workbook, args.sheet, args.metric)
    model_styles = _models_to_plot(args.exclude_no_decay)

    written: list[Path] = []
    missing: list[tuple[str, str]] = []
    for deal in args.deals:
        for scenario in args.scenarios:
            output_path = plot_deal_scenario(
                records,
                deal,
                scenario,
                args.output_dir,
                args.max_month,
                model_styles,
                args.metric,
            )
            if output_path is None:
                missing.append((deal, scenario))
            else:
                written.append(output_path)

    print(f"Wrote {len(written)} plots to {args.output_dir}")
    for path in written:
        print(path)
    if missing:
        print("Missing deal/scenario combinations:")
        for deal, scenario in missing:
            print(f"- {deal}: {scenario}")

    return 0 if written else 1


if __name__ == "__main__":
    raise SystemExit(main())
